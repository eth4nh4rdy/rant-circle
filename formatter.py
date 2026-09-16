from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt, RGBColor

SPACE_INDIGO = RGBColor(0x28, 0x2C, 0x45)
GREY = RGBColor(0x80, 0x80, 0x80)
BLACK = RGBColor(0x00, 0x00, 0x00)

A4_WIDTH = Mm(210)
A4_HEIGHT = Mm(297)


def create_document():
    document = Document()

    section = document.sections[0]
    section.page_width = A4_WIDTH
    section.page_height = A4_HEIGHT

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    # Ensure Arial is used for East Asian text runs too, not just Latin.
    rpr = normal_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Arial")

    return document


def add_header(document, name_field=True, date=None, topic=None):
    title = document.add_paragraph()
    title_run = title.add_run("Rant Circle")
    title_run.font.size = Pt(34)
    title_run.font.bold = True
    title_run.font.color.rgb = SPACE_INDIGO

    if name_field:
        name_paragraph = document.add_paragraph()
        name_run = name_paragraph.add_run("Name: ___________________")
        name_run.font.color.rgb = BLACK

    date_paragraph = document.add_paragraph()
    if date:
        formatted_date = datetime.strptime(date, "%Y-%m-%d").strftime("%B %-d, %Y")
    else:
        formatted_date = datetime.today().strftime("%B %-d, %Y")
    date_run = date_paragraph.add_run(f"Date: {formatted_date}")
    date_run.font.color.rgb = BLACK

    topic_paragraph = document.add_paragraph()
    topic_run = topic_paragraph.add_run(f"Topic: {topic if topic else 'General'}")
    topic_run.font.color.rgb = BLACK


def add_footer(document):
    footer = document.sections[0].footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run("www.primo-english.xyz")
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = GREY

    page_number_paragraph = footer.add_paragraph()
    page_number_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    run = page_number_paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr_text)
    run._r.append(fld_end)


def add_section_divider(document, text=None):
    if text:
        label_paragraph = document.add_paragraph()
        label_run = label_paragraph.add_run(text)
        label_run.font.size = Pt(9)
        label_run.font.bold = True
        label_run.font.color.rgb = SPACE_INDIGO

    divider_paragraph = document.add_paragraph()
    p_pr = divider_paragraph._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom_border = OxmlElement("w:bottom")
    bottom_border.set(qn("w:val"), "single")
    bottom_border.set(qn("w:sz"), "6")
    bottom_border.set(qn("w:space"), "1")
    bottom_border.set(qn("w:color"), "282C45")
    p_borders.append(bottom_border)
    p_pr.append(p_borders)


if __name__ == "__main__":
    document = create_document()
    add_footer(document)
    add_header(document, name_field=True, date="2026-09-16", topic="Work Hours")
    document.add_paragraph("This is a test of the document skeleton.")
    add_section_divider(document, text="Test Divider")
    document.save("output/skeleton_test.docx")
    print("Saved output/skeleton_test.docx")
